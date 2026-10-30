"""Real-format conversion, authenticated preview and byte-exact student download."""
import os,sys,json,time,subprocess,hashlib,io,tempfile,urllib.request,urllib.error,unittest
from pathlib import Path
from PIL import Image,ImageDraw
ROOT=Path(__file__).resolve().parents[1]
def fixtures(folder):
 from docx import Document
 from pptx import Presentation
 import pymupdf
 folder=Path(folder);folder.mkdir(parents=True,exist_ok=True)
 im=Image.new('RGB',(900,600),'#dcece3');ImageDraw.Draw(im).text((60,100),'Knowledge map - image format preview',fill='black')
 (folder/'diagram.svg').write_text('<svg xmlns="http://www.w3.org/2000/svg" width="1000" height="600"><rect x="20" y="20" width="960" height="560" rx="20" fill="#e7f0df"/><path d="M80 300 L800 180" stroke="green" stroke-width="5"/><text x="80" y="110" font-size="38">SVG vector preview</text><text x="80" y="180" font-size="36">环境评价知识图谱</text></svg>')
 for ext in ['png','jpg','webp','gif','bmp','tiff']:im.save(folder/('image.'+ext))
 doc=Document();doc.add_heading('Knowledge map - Word preview',0);doc.add_paragraph('Environment assessment: identify, predict, evaluate, mitigate.');doc.add_page_break();doc.add_heading('Second page',1);doc.add_paragraph('Page navigation test');doc.save(folder/'document.docx')
 prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[1]);slide.shapes.title.text='Knowledge map - PowerPoint preview';slide.placeholders[1].text='Identify\nPredict\nEvaluate\nMitigate';slide=prs.slides.add_slide(prs.slide_layouts[1]);slide.shapes.title.text='Second slide';slide.placeholders[1].text='Page navigation test';prs.save(folder/'slides.pptx')
 pdf=pymupdf.open();page=pdf.new_page();page.insert_text((60,90),'Knowledge map - PDF preview');page=pdf.new_page();page.insert_text((60,90),'Second page - navigation test');pdf.save(folder/'document.pdf');pdf.close()
 exe='/Applications/LibreOffice.app/Contents/MacOS/soffice' if sys.platform=='darwin' else 'libreoffice'
 for source,fmt in [('document.docx','doc:MS Word 97'),('slides.pptx','ppt:MS PowerPoint 97')]:
  with tempfile.TemporaryDirectory() as profile:
   subprocess.run([exe,'-env:UserInstallation='+Path(profile).as_uri(),'--headless','--convert-to',fmt,'--outdir',str(folder),str(folder/source)],capture_output=True,check=True,timeout=90)
 return sorted(p for p in folder.iterdir() if p.suffix in ['.svg','.png','.jpg','.webp','.gif','.bmp','.tiff','.pdf','.doc','.docx','.ppt','.pptx'])
class Formats(unittest.TestCase):
 def test_all_formats(self):
  from test_api import port
  with tempfile.TemporaryDirectory() as temp:
   files=fixtures(Path(temp)/'fixtures');self.assertEqual(len(files),12)
   p=port();base=f'http://127.0.0.1:{p}';env={**os.environ,'KM_PORT':str(p),'KM_DEMO':'1','KM_DATA':str(Path(temp)/'data')};proc=subprocess.Popen([sys.executable,str(ROOT/'server.py')],env=env)
   def req(path,token='demo-teacher',body=None,method='GET'):
    with urllib.request.urlopen(urllib.request.Request(base+'/api/'+path,data=body,method=method,headers={'Authorization':'Bearer '+token}),timeout=180) as r:return r.headers.get('Content-Type'),r.read()
   try:
    for _ in range(100):
     try:urllib.request.urlopen(base+'/api/config');break
     except OSError:time.sleep(.05)
    report=[]
    for f in files:
     raw=f.read_bytes();_,response=req('maps?'+urllib.parse.urlencode({'name':f.name,'title':f.name,'classId':'demo-class'}),body=raw,method='POST');m=json.loads(response);self.assertEqual(m['mode'],'file');self.assertFalse(m['previewError'],m['previewError'])
     with self.assertRaises(urllib.error.HTTPError) as e:req('maps/'+m['id']+'/preview','demo-student')
     self.assertEqual(e.exception.code,404);e.exception.close()
     _,response=req('maps/'+m['id'],body=json.dumps({**m,'state':'published'}).encode(),method='PUT');m=json.loads(response)
     mime,preview=req('maps/'+m['id']+'/preview','demo-student')
     if mime=='application/pdf':
      import pymupdf
      with pymupdf.open(stream=preview,filetype='pdf') as pdf:self.assertGreater(len(pdf),0);self.assertIn('preview',''.join(p.get_text() for p in pdf).lower());pix=pdf[0].get_pixmap();self.assertGreater(pix.width,0)
      self.assertEqual(m['pageCount'],2)
      _,rendered=req('maps/'+m['id']+'/page/2','demo-student')
      with Image.open(io.BytesIO(rendered)) as image:image.load();self.assertGreater(image.width,0)
     elif mime=='image/svg+xml':self.assertEqual(preview,raw)
     else:
      with Image.open(io.BytesIO(preview)) as image:image.load();self.assertGreater(image.width,0)
     _,pdfdata=req('maps/'+m['id']+'/pdf','demo-student')
     import pymupdf
     with pymupdf.open(stream=pdfdata,filetype='pdf') as exported:
      self.assertGreater(len(exported),0);exported[0].get_pixmap()
      if f.suffix=='.svg':self.assertEqual(exported[0].get_images(),[]);self.assertGreater(len(exported[0].get_drawings()),0);self.assertIn('环境评价',exported[0].get_text())
     if f.suffix=='.pdf':self.assertEqual(pdfdata,raw)
     _,cached=req('maps/'+m['id']+'/pdf','demo-student');self.assertEqual(cached,pdfdata)
     _,download=req('maps/'+m['id']+'/file','demo-student');self.assertEqual(download,raw)
     with self.assertRaises(urllib.error.HTTPError) as e:req('maps/'+m['id']+'/preview','demo-other')
     self.assertEqual(e.exception.code,404);e.exception.close()
     with self.assertRaises(urllib.error.HTTPError) as e:req('maps/'+m['id']+'/pdf','demo-other')
     self.assertEqual(e.exception.code,404);e.exception.close()
     report.append({'format':f.suffix,'preview':mime,'download_sha256':hashlib.sha256(download).hexdigest(),'pdf_export_valid':True,'pdf_cache_consistent':True,'draft_hidden':True,'other_class_denied':True});print('PASS',f.suffix,mime,flush=True)
    (ROOT/'work/format-verification.json').write_text(json.dumps(report,indent=2))
   finally:proc.terminate();proc.wait()
if __name__=='__main__':unittest.main()
